import openpyxl

from csmmap.excel import detect_columns, load_folder


def test_load_fixture(folder):
    ds = load_folder(folder)
    assert len(ds.stores) == 12
    assert len(ds.dealers) == 5
    assert ds.files[0].name == "CSM.xlsx" and ds.files[0].sheet == "CSM"
    by_name = {d.name: d for d in ds.dealers}
    assert by_name["(주)한빛유통"].address == "경기도 부천시 원미구 길주로 210"
    assert by_name["한빛유통"].address == ""
    # 끝 공백 정리
    assert any(s.name == "샘플마트 양천점" for s in ds.stores)
    # 경고: 주소 없는 거래처 + 유사 명칭
    assert any("주소 없음" in w for w in ds.warnings)
    assert any("유사 대리점명" in w for w in ds.warnings)
    assert by_name["(주)한빛유통"].similar_to == ["한빛유통"]


def test_color_idx_follows_store_count(folder):
    ds = load_folder(folder)
    counts = [len(d.store_ids) for d in ds.dealers]
    assert counts == sorted(counts, reverse=True)
    assert [d.color_idx for d in ds.dealers] == list(range(len(ds.dealers)))


def test_ignores_temp_and_merges_multiple_files(folder):
    (folder / "~$CSM.xlsx").write_bytes(b"junk")
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B1"], ws["C1"], ws["D1"] = "대리점명", "거래처명", "주소"
    ws["B2"], ws["C2"], ws["D2"] = "새대리점", "추가마트", "서울특별시 송파구 올림픽로 300"
    wb.save(folder / "extra.xlsx")
    ds = load_folder(folder)
    assert len(ds.files) == 2
    assert len(ds.stores) == 13
    assert any(d.name == "새대리점" for d in ds.dealers)


def test_detect_columns_moved_headers(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A2"], ws["B2"], ws["C2"], ws["F2"], ws["G2"] = "대리점명", "거래처명", "주소", "대리점명", "주소"
    wb.save(tmp_path / "x.xlsx")
    ws2 = openpyxl.load_workbook(tmp_path / "x.xlsx").active
    cols = detect_columns(ws2)
    assert (cols.store_dealer, cols.store_name, cols.store_addr) == (1, 2, 3)
    assert (cols.dealer_name, cols.dealer_addr, cols.header_row) == (6, 7, 2)
