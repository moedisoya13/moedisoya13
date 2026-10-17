"""폴더에 놓인 적대적 파일(조작된 캐시·이상한 xlsx)과 내보내기 탈출 시도."""
import json
import threading
import urllib.request

import openpyxl

from csmmap import model
from csmmap.config import Config
from csmmap.export import _safe_json, build_html
from csmmap.geocode import Cache, Geocoder, GeoResult
from csmmap.pipeline import Pipeline
from csmmap.server import serve


def test_tampered_cache_entries_are_ignored(tmp_path, regions):
    path = tmp_path / "geocode_cache.json"
    path.write_text(json.dumps({
        "xss": {"lat": 37.5, "lng": 127.0, "precision": "<img src=x onerror=alert(1)>", "provider": "", "query": "", "ts": 0},
        "unknown": {"lat": 37.5, "lng": 127.0, "precision": "exact", "evil": 1},
        "nan": {"lat": "nan", "lng": 127.0, "precision": "exact", "provider": "", "query": "", "ts": 0},
        "range": {"lat": 91, "lng": 127.0, "precision": "exact", "provider": "", "query": "", "ts": 0},
        "notdict": 5,
        "good": {"lat": 37.5, "lng": 127.0, "precision": "exact", "provider": "kakao:address", "query": "q", "ts": 0},
    }), encoding="utf-8")
    c = Cache(path)
    assert c.get("xss") is None and c.get("nan") is None
    assert c.get("unknown").precision == model.EXACT   # 모르는 키는 버리고 나머지는 쓴다
    assert c.get("range") is None and c.get("notdict") is None
    assert c.get("good").precision == model.EXACT


def test_corrupt_cache_file_does_not_break_startup(folder, regions):
    (folder / "geocode_cache.json").write_text("[1,2,3", encoding="utf-8")
    p = Pipeline(folder, regions, Config(folder), demo=True)
    p.apply_cached()
    assert len(p.dataset.stores) == 12


def test_kakao_nonfinite_coordinates_are_rejected(regions, tmp_path):
    class Fake:
        def address(self, q): return {"x": "1e999", "y": "37.5"}
    g = Geocoder(regions, Cache(tmp_path / "c.json"), Fake())
    r = g.geocode_one("경기도 부천시 소사구 범안로 103", "가게")
    assert r.precision == model.NEEDS_FIX  # 비정상 좌표는 적중으로 치지 않는다


def test_safe_json_blocks_script_data_escape():
    out = _safe_json({"name": "<!--<script>", "b": "</script><b>"})
    assert "<" not in out and "\\u003c!--\\u003cscript>" in out
    assert json.loads(out)["name"] == "<!--<script>"


def test_export_html_has_no_font_link_and_no_raw_lt_in_data(folder, regions):
    wb = openpyxl.load_workbook(folder / "CSM.xlsx"); ws = wb.active
    ws["C2"] = "<!--<script>악성"
    wb.save(folder / "CSM.xlsx")
    p = Pipeline(folder, regions, Config(folder), demo=True)
    p.geocode()
    html = build_html(p.payload(for_export=True))
    assert "fonts.googleapis.com" not in html
    data_line = next(l for l in html.splitlines() if l.startswith("window.CSM_DATA"))
    assert "<" not in data_line[len("window.CSM_DATA = "):]


def test_chartsheet_only_workbook_is_skipped(folder):
    from csmmap.excel import load_folder
    wb = openpyxl.Workbook(); wb.remove(wb.active); wb.create_chartsheet("차트")
    wb.save(folder / "charts.xlsx")
    ds = load_folder(folder)
    assert len(ds.stores) == 12
    assert any("charts.xlsx" in w and ("워크시트가 없음" in w or "열 수 없음" in w) for w in ds.warnings)


def test_negative_content_length_rejected(folder, regions):
    p = Pipeline(folder, regions, Config(folder), demo=True)
    httpd, state, port = serve(p, port=0, auto_exit=False)
    threading.Thread(target=httpd.serve_forever, kwargs={"poll_interval": 0.1}, daemon=True).start()
    try:
        req = urllib.request.Request(f"http://127.0.0.1:{port}/api/scan", data=b"", method="POST",
                                     headers={"X-CSM-Token": state.token, "Content-Length": "-1"})
        try:
            urllib.request.urlopen(req, timeout=5)
            assert False
        except urllib.error.HTTPError as e:
            assert e.code == 400
    finally:
        httpd.shutdown(); httpd.server_close()


def test_state_survives_nan_in_dataset(folder, regions):
    """직렬화 불가 값이 끼어도 연결을 끊지 않고 500 JSON 으로 답한다."""
    p = Pipeline(folder, regions, Config(folder), demo=True)
    p.dataset.stores[0].lat = float("nan")
    httpd, state, port = serve(p, port=0, auto_exit=False)
    threading.Thread(target=httpd.serve_forever, kwargs={"poll_interval": 0.1}, daemon=True).start()
    try:
        req = urllib.request.Request(f"http://127.0.0.1:{port}/api/state", headers={"X-CSM-Token": state.token})
        try:
            urllib.request.urlopen(req, timeout=5)
            assert False
        except urllib.error.HTTPError as e:
            assert e.code == 500 and "error" in json.loads(e.read())
    finally:
        httpd.shutdown(); httpd.server_close()
