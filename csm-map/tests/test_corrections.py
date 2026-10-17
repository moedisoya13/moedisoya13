"""정정 요청 목록 xlsx."""
import json
import threading
import urllib.error
import urllib.request

import openpyxl
import pytest

from csmmap import model
from csmmap.config import Config
from csmmap.export import CORRECTION_COLUMNS, FIX_COL, correction_rows, export_corrections
from csmmap.geocode import Cache, Geocoder
from csmmap.pipeline import Pipeline
from csmmap.server import serve


class NoHit:
    """카카오가 어떤 주소도 못 찾는 상황."""
    def address(self, q):
        return None


def live_pipeline(folder, regions, monkeypatch):
    """키가 있는 것처럼, 그러나 모든 주소 검색이 실패하도록 만든 파이프라인."""
    cfg = Config(folder)
    cfg.data["api_key"] = "testkey123"
    p = Pipeline(folder, regions, cfg, demo=False)
    monkeypatch.setattr(p, "make_geocoder", lambda: Geocoder(regions, p.cache, NoHit(), overrides=p.overrides.data))
    p.geocode()
    return p


def test_all_misses_become_needs_fix(folder, regions, monkeypatch):
    p = live_pipeline(folder, regions, monkeypatch)
    counts = p.payload()["meta"]["counts"]
    assert counts["needs_fix"] == 12 and counts["exact"] == 0 and counts["approx"] == 0
    d = {x.name: x for x in p.dataset.dealers}
    assert d["(주)한빛유통"].precision == model.NEEDS_FIX
    assert d["한빛유통"].precision == model.FAILED  # 대리점 주소가 비면 정정 대상 아님(선택 입력)


def test_correction_rows_and_order(folder, regions, monkeypatch):
    p = live_pipeline(folder, regions, monkeypatch)
    rows = correction_rows(p)
    assert len(rows) == 12 + 4                       # 거래처 12 + 주소 있는 대리점 4
    assert [r[0] for r in rows] == list(range(1, 17))
    assert rows[0][1] == "거래처" and rows[-1][1] == "대리점"
    empty = next(r for r in rows if r[3] == "샘플마트 주소없음")
    assert empty[5] == "주소 없음"


def test_export_xlsx_layout_and_formula_guard(folder, regions, monkeypatch):
    wb = openpyxl.load_workbook(folder / "CSM.xlsx"); ws = wb.active
    ws["D2"] = '=HYPERLINK("http://evil.example","클릭")'
    ws["D2"].data_type = "s"                         # 원본에 문자열로 들어 있는 경우
    wb.save(folder / "CSM.xlsx")
    p = live_pipeline(folder, regions, monkeypatch)
    out, n = export_corrections(p)
    assert out.name.startswith("CSM_정정요청_") and out.suffix == ".xlsx" and n == 16
    book = openpyxl.load_workbook(out)
    assert book.sheetnames == ["정정 요청", "안내"]
    sh = book["정정 요청"]
    assert [c.value for c in sh[1]] == [t for t, _ in CORRECTION_COLUMNS]
    assert sh.cell(row=1, column=FIX_COL).value == "수정 주소"
    assert sh.freeze_panes == "E2" and sh.auto_filter.ref.startswith("A1:")
    cells = [sh.cell(row=r, column=5) for r in range(2, sh.max_row + 1)]
    evil = next(c for c in cells if str(c.value).startswith("=HYPERLINK"))
    assert evil.data_type == "s"                     # 수식이 아니라 문자열


def test_no_targets_raises(folder, regions):
    p = Pipeline(folder, regions, Config(folder), demo=True)
    p.geocode()
    for s in p.dataset.stores:
        if s.precision == model.NEEDS_FIX:
            s.precision = model.APPROX
    with pytest.raises(ValueError):
        export_corrections(p)


def test_server_corrections_endpoint(folder, regions, monkeypatch):
    p = live_pipeline(folder, regions, monkeypatch)
    httpd, state, port = serve(p, port=0, auto_exit=False)
    threading.Thread(target=httpd.serve_forever, kwargs={"poll_interval": 0.1}, daemon=True).start()
    try:
        def post(body):
            req = urllib.request.Request(f"http://127.0.0.1:{port}/api/export", data=json.dumps(body).encode(), method="POST",
                                         headers={"Content-Type": "application/json", "X-CSM-Token": state.token})
            try:
                with urllib.request.urlopen(req, timeout=5) as r:
                    return r.status, json.loads(r.read())
            except urllib.error.HTTPError as e:
                return e.code, json.loads(e.read())
        st, r = post({"format": "corrections"})
        assert st == 200 and r["count"] == 16 and r["name"].startswith("CSM_정정요청_")
        for s in p.dataset.stores + p.dataset.dealers:
            s.precision = model.EXACT
        st, r = post({"format": "corrections"})
        assert st == 400 and "정정할 항목" in r["error"]
    finally:
        httpd.shutdown(); httpd.server_close()


def test_cached_needs_fix_shows_on_restart(folder, regions, monkeypatch):
    live_pipeline(folder, regions, monkeypatch)       # 캐시에 needs_fix 저장
    p2 = Pipeline(folder, regions, Config(folder), demo=True)
    p2.apply_cached()                                  # 네트워크 없이 즉시
    assert sum(s.precision == model.NEEDS_FIX for s in p2.dataset.stores) == 12
    assert Cache(folder / "geocode_cache.json").get("경기도 시흥시 배곧4로 32-32").reason == "주소 검색 결과 없음"


def test_own_correction_file_is_not_read_back_as_data(folder, regions, monkeypatch):
    p = live_pipeline(folder, regions, monkeypatch)
    out, _ = export_corrections(p)
    from csmmap.excel import load_folder
    assert len(load_folder(folder).stores) == 12          # 같은 폴더에 있어도 읽지 않음
    out.rename(folder / "정정회신.xlsx")                    # 이름을 바꿔도 문서 속성으로 알아본다
    ds = load_folder(folder)
    assert len(ds.stores) == 12 and len(ds.files) == 1
